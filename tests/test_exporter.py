from exporter import export_csv, export_excel
from openpyxl import load_workbook

def test_export_csv_union_and_formula_safety(tmp_path):
    path = tmp_path / "rows.csv"
    export_csv([{"a": "=1+1"}, {"a": "ok", "b": "x"}], str(path))
    text = path.read_text(encoding="utf-8-sig")
    assert "'=1+1" in text and "b" in text

def test_export_excel_preserves_types(tmp_path):
    path = tmp_path / "rows.xlsx"
    export_excel([{"n": 123, "f": 4.5, "z": None, "s": "a\x0bb"}], str(path))
    ws = load_workbook(path).active
    assert ws["A2"].value == 123 and ws["B2"].value == 4.5 and ws["C2"].value is None and ws["D2"].value == "ab"
